import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

EXCEL_FILE = "users.xlsx"


def save_user_to_excel(fullname, email, password):
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet(title="Users")
        else:
            ws.title = "Users"
        ws.append(["Full Name", "Email", "Password", "Registration Date"])
        wb.save(EXCEL_FILE)

    wb = load_workbook(EXCEL_FILE)
    ws = wb["Users"]
    
    registration_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws.append([fullname, email, password, registration_date])
    
    wb.save(EXCEL_FILE)


def authenticate_user(username_or_email, password):
    if not os.path.exists(EXCEL_FILE):
        return False, "No registered users found. Please sign up first."

    wb = load_workbook(EXCEL_FILE)
    if "Users" not in wb.sheetnames:
        return False, "No registered users found. Please sign up first."
    
    ws = wb["Users"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 3:
            continue
        
        db_fullname, db_email, db_password = str(row[0]), str(row[1]), str(row[2])

        if username_or_email.lower() in (db_fullname.lower(), db_email.lower()):
            if password == db_password:
                return True, db_fullname  
            else:
                return False, "Incorrect password."

    return False, "User account not found."